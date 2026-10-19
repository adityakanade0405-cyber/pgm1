import json
import os
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from dotenv import load_dotenv

load_dotenv()

OUTPUT_DIR = Path("output")
OUTPUT_DIR.mkdir(exist_ok=True)


def get_ai_content(topic):
    """Generate topic-specific presentation content using the OpenAI API."""
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        return None

    try:
        from openai import OpenAI
        client = OpenAI(api_key=api_key)

        prompt = f"""
Create accurate, beginner-friendly presentation content about: {topic}

Return ONLY valid JSON with this exact structure:
{{
  "title": "Presentation title",
  "objectives": ["objective 1", "objective 2", "objective 3", "objective 4"],
  "slides": [
    {{
      "heading": "Slide heading",
      "bullets": ["bullet 1", "bullet 2", "bullet 3", "bullet 4"],
      "speaker_notes": "A short explanation the presenter can say aloud."
    }}
  ]
}}

Rules:
- Include exactly 6 content slides in the slides array.
- Use a logical sequence: Introduction, Key Concepts/Features, How It Works,
  Applications, Benefits and Challenges, Future Scope/Conclusion.
- Make all content specific to the requested topic, not generic filler.
- Use clear, simple English and factual claims.
- Keep each bullet concise and provide useful speaker notes for every slide.
- Do not include markdown fences or any text outside the JSON.
"""
        response = client.responses.create(
            model=os.getenv("OPENAI_MODEL", "gpt-4.1-mini"),
            input=prompt
        )
        raw = response.output_text.strip()
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)
        data = json.loads(raw)

        if not isinstance(data.get("slides"), list) or len(data["slides"]) < 6:
            raise ValueError("The AI did not return at least six content slides.")
        return data
    except Exception as exc:
        print(f"\nAI generation could not be completed: {exc}")
        print("Switching to built-in template content so you can still create a PPT.")
        return None


def get_template_content(topic):
    """Offline fallback. This uses a template, not a live AI model."""
    return {
        "title": f"Introduction to {topic}",
        "objectives": [
            f"Understand the meaning and purpose of {topic}",
            f"Identify important concepts related to {topic}",
            f"Explore possible uses and benefits of {topic}",
            f"Discuss challenges and future developments"
        ],
        "slides": [
            {
                "heading": "Introduction",
                "bullets": [
                    f"Topic: {topic}",
                    "Meaning and basic overview",
                    "Why the topic matters",
                    "Where it may be used"
                ],
                "speaker_notes": f"Introduce {topic}. Explain its meaning, basic purpose, and why people study or use it."
            },
            {
                "heading": "Key Concepts and Features",
                "bullets": [
                    f"Main ideas related to {topic}",
                    "Important terms and components",
                    "Common features or characteristics",
                    "Examples can help explain the concepts"
                ],
                "speaker_notes": f"Describe the main concepts and features of {topic}. Add accurate examples that you have verified."
            },
            {
                "heading": "How It Works",
                "bullets": [
                    "Identify the input or requirements",
                    "Explain the main steps or process",
                    "Describe the output or result",
                    "Mention tools or components involved"
                ],
                "speaker_notes": f"Explain the process behind {topic} in a simple step-by-step way. Adapt these points to the actual topic."
            },
            {
                "heading": "Applications",
                "bullets": [
                    "Education and learning",
                    "Business and industry",
                    "Research and development",
                    "Everyday digital services, where relevant"
                ],
                "speaker_notes": f"Discuss real applications of {topic}. Replace any example that is not relevant to the selected topic."
            },
            {
                "heading": "Benefits and Challenges",
                "bullets": [
                    "Potential to improve efficiency",
                    "Can help solve specific problems",
                    "May require training, time, or money",
                    "Privacy, security, or reliability may matter"
                ],
                "speaker_notes": f"Present both the benefits and limitations of {topic}. Avoid claiming that every benefit applies in every situation."
            },
            {
                "heading": "Future Scope and Conclusion",
                "bullets": [
                    "Technology and practices may continue to evolve",
                    "New use cases may emerge",
                    "Responsible use is important",
                    "Summarize the main learning points"
                ],
                "speaker_notes": f"Summarize the presentation about {topic}. Explain that future developments depend on research, practical needs, and responsible implementation."
            }
        ]
    }


def add_bullet_slide(prs, heading, bullets, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = heading

    body = slide.placeholders[1].text_frame
    body.clear()
    body.word_wrap = True

    for index, bullet in enumerate(bullets):
        paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
        paragraph.text = str(bullet)
        paragraph.level = 0
        paragraph.font.size = Pt(22)
        paragraph.font.color.rgb = RGBColor(45, 55, 72)
        paragraph.space_after = Pt(10)

    try:
        slide.notes_slide.notes_text_frame.text = notes
    except Exception:
        print(f"Warning: speaker notes could not be added to slide '{heading}'.")

    return slide


def create_presentation(topic, data):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = data.get("title", topic)
    slide.placeholders[1].text = "AI-Powered Presentation Content Generator"

    # Objectives slide
    add_bullet_slide(
        prs,
        "Objectives",
        data.get("objectives", [
            f"Understand {topic}",
            "Learn key concepts",
            "Explore applications",
            "Review benefits and challenges"
        ]),
        f"Explain what the presentation aims to teach about {topic}. Briefly introduce each objective."
    )

    # Six or more content slides
    for item in data["slides"]:
        heading = item.get("heading", "Content")
        bullets = item.get("bullets", [])
        notes = item.get("speaker_notes", f"Explain the main points about {heading}.")
        add_bullet_slide(prs, heading, bullets, notes)

    safe_topic = re.sub(r"[^A-Za-z0-9_-]+", "_", topic).strip("_")[:50] or "Presentation"
    output_path = OUTPUT_DIR / f"{safe_topic}_Presentation.pptx"
    prs.save(output_path)
    return output_path, len(prs.slides)


def main():
    print("=" * 55)
    print(" AI-POWERED PRESENTATION CONTENT GENERATOR")
    print("=" * 55)
    topic = input("Enter a computer/IT/emerging technology topic: ").strip()

    if not topic:
        print("Please enter a topic and run the program again.")
        return

    print("\nGenerating presentation content...")
    data = get_ai_content(topic)

    if data is None:
        data = get_template_content(topic)
        print("Note: Template mode was used. Add an API key to enable AI-generated content.")

    try:
        path, slide_count = create_presentation(topic, data)
        print("\nSUCCESS: Presentation created!")
        print(f"File: {path.resolve()}")
        print(f"Total slides: {slide_count}")
        print("Content slides include speaker notes.")
        print("\nPlease open the PPT and review factual accuracy and readability.")
    except Exception as exc:
        print(f"\nCould not create the presentation: {exc}")


if __name__ == "__main__":
    main()
